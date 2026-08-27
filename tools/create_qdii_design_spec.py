from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "QDII直销代销额度卡片设计规范.docx"
PREVIEW = ROOT / "data/cards/xiaohongshu-direct-sales/qdii-主动型-direct-sales-20260827-header-preview.png"

NAVY = "05060B"
GOLD = "D4B26A"
PALE_GOLD = "F2E9D2"
RED = "EF3B3B"
INK = "E7ECF3"
MUTED = "8A93A6"
BLUE = "141A2E"
DOC_INK = RGBColor(30, 35, 48)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def borders(cell, color="D9DEE8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        node = tc_borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def fmt_run(run, size=10.5, color=DOC_INK, bold=False, font="Hiragino Sans GB"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_text(doc, text, *, size=10.5, color=DOC_INK, bold=False, align=None, before=0, after=6, font="Hiragino Sans GB"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    fmt_run(p.add_run(text), size=size, color=color, bold=bold, font=font)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    fmt_run(p.add_run(text), size=16 if level == 1 else 12.5, color=RGBColor(25, 63, 105), bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.375)
    p.paragraph_format.first_line_indent = Inches(-.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    fmt_run(p.add_run("• "), size=10.5, color=RGBColor(25, 63, 105), bold=True)
    fmt_run(p.add_run(text), size=10.5)


def fixed_table(doc, rows, widths, header=True):
    table = doc.add_table(rows=0, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_margins(cells[i])
            borders(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if header and r_i == 0:
                shade(cells[i], "E8EEF5")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            fmt_run(p.add_run(value), size=9.5, bold=(header and r_i == 0))
    return table


def set_header_footer(section):
    section.header_distance = Inches(.492)
    section.footer_distance = Inches(.492)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_run(hp.add_run("QDII QUOTA RADAR  ·  DESIGN SPEC"), size=8, color=RGBColor(120, 128, 143), font="Arial")
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_run(fp.add_run("内部设计规范 · 2026-08-27"), size=8, color=RGBColor(120, 128, 143), font="Arial")


def image_font(size, bold=False):
    path = "/System/Library/Fonts/Supplemental/Songti.ttc"
    return ImageFont.truetype(path, size, index=1 if bold else 0)


def make_image_page(path, title, subtitle, sections):
    canvas = Image.new("RGB", (1500, 1950), "#FFFFFF")
    d = ImageDraw.Draw(canvas)
    serif = image_font(36)
    strong = image_font(48, True)
    body = image_font(28)
    small = image_font(23)
    d.text((112, 90), "QDII QUOTA RADAR · DESIGN SPEC", font=small, fill="#8A93A6")
    d.text((112, 165), title, font=strong, fill="#172D51")
    d.text((114, 237), subtitle, font=serif, fill="#58677A")
    d.line((112, 300, 1388, 300), fill="#D4B26A", width=3)
    y = 355
    for heading, lines in sections:
        d.text((112, y), heading, font=image_font(34, True), fill="#173F69")
        y += 58
        for line in lines:
            # A compact, deterministic line wrapper for CJK documentation text.
            chunks = [line[i:i+32] for i in range(0, len(line), 32)]
            for chunk in chunks:
                d.text((142, y), chunk, font=body, fill="#273445")
                y += 47
        y += 22
    d.text((112, 1875), "内部设计规范 · 2026-08-27", font=small, fill="#8A93A6")
    canvas.save(path, "PNG")


def create_image_backed_doc():
    """Use rasterized Chinese pages because the headless Office renderer lacks CJK glyph support."""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    asset_dir = ROOT / "docs" / "assets"
    asset_dir.mkdir(parents=True, exist_ok=True)
    pages = [
        ("01  内容与数据规则", "一张卡片只呈现一个基金类别；主动型与被动型分别输出。", [
            ("固定信息", ["眉题：ACTIVE / PASSIVE QDII · RETURN & QUOTA", "标题：QDII 主动型基金直销 / 代销额度（或被动型）", "日期：YYYY — MM — DD", "列：基金名称｜近一年收益率｜直销额度｜代销额度"]),
            ("数据处理", ["数据源：config.json 定义基金池与显示名；行情与额度沿用抓取逻辑。", "直销额度仅使用基金公司官网当前可验证数据；状态非 ok 不进入卡片。", "未实现直销额度、不支持或抓取失败的基金，先不获取、也不展示。", "每个类别按近一年收益率由高到低排序；并列时按基金代码升序。"]),
            ("已知抓取前排除", ["016702、017093、018036、019442。输出与排除清单保留在 manifest 中。"]),
        ]),
        ("02  版式与信息层级", "采用已确认的开放式表头：连续表格，不使用卡片式表头或单基金方块。", [
            ("画布与标题区", ["画布：1080 × 1440 px 竖版；四角为金色 L 形细线。", "眉题 y=112；主标题 y=156；日期 y=228；短金线 y=282。", "主标题保持常规宋体字重，不额外加粗。"]),
            ("开放式表头", ["表头从 y=330 开始：仅文字与一条金色分隔线，不填充、不圆角。", "列锚点：基金名称 x=82 左对齐；收益率 x=620、直销 x=840、代销 x=998 右对齐。", "数据行 60 px 高；行间用低透明度细线；基金之间不得产生间隔卡片。"]),
            ("字段细节", ["基金名称过长时省略；基金代码作为次级文字紧随其后。", "额度用紧凑写法，如“1万”；不使用序号列与底部注释。"]),
        ]),
        ("03  字体、颜色与状态", "方案 C：宋体中文 + 衬线数字，控制为深色、金色、红色的低干扰体系。", [
            ("字体", ["中文标题、基金名称、中文额度状态：Songti SC（宋体）。", "百分比、纯数字、日期：DM Serif Display。", "英文眉题：Cormorant Garamond。"]),
            ("颜色", ["主背景 #05060B；主金色 #D4B26A；浅金文字/金额 #F2E9D2。", "正收益与“暂停” #EF3B3B；基金名称 #E7ECF3；基金代码 #8A93A6。", "#141A2E 仅作为旧版兼容色保留；当前开放式表头不填充深蓝。"]),
            ("状态与验收", ["“暂停”一律红色，其他额度一律浅金；不引入绿色。", "输出前确认：标题类型正确、数值右对齐、无表头胶囊、无逐项卡片间隔。", "最终各输出一张主动型与被动型 1080×1440 PNG。"]),
        ]),
    ]
    page_paths = []
    for i, (title, subtitle, sections) in enumerate(pages, 1):
        page = asset_dir / f"qdii-design-spec-{i}.png"
        make_image_page(page, title, subtitle, sections)
        page_paths.append(page)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.5)
    section.left_margin = section.right_margin = Inches(.5)
    for i, page in enumerate(page_paths):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(page), width=Inches(7.35))
        if i < len(page_paths) - 1:
            doc.add_page_break()
    doc.core_properties.title = "QDII 直销 / 代销额度卡片设计规范"
    doc.core_properties.subject = "小红书基金额度卡片的视觉与数据规则"
    doc.core_properties.author = "QDII Quota Radar"
    doc.save(OUT)
    print(OUT)


def main():
    create_image_backed_doc()
    return
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    set_header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Hiragino Sans GB"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal.font.size = Pt(11)

    add_text(doc, "QDII QUOTA RADAR", size=10, color=RGBColor(159, 120, 29), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=68, after=12, font="Arial")
    add_text(doc, "QDII 直销 / 代销额度卡片", size=27, color=RGBColor(18, 30, 51), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_text(doc, "设计与数据特征规范 · 已确认版本", size=13, color=RGBColor(70, 83, 105), align=WD_ALIGN_PARAGRAPH.CENTER, after=22)
    add_text(doc, "用途：用于按近一年收益率排序呈现 QDII 主动型与被动型基金的直销、代销额度。\n版本基线：无卡片式表头、连续表格行、宋体中文方案（C）。", size=10.5, color=RGBColor(70, 83, 105), align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    if PREVIEW.exists():
        doc.add_picture(str(PREVIEW), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(doc, "已确认视觉基线：开放式表头 + 一条金色分隔线；不使用独立表头卡片。", size=8.5, color=RGBColor(95, 105, 120), align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    doc.add_page_break()
    add_heading(doc, "1. 内容结构与数据规则")
    add_text(doc, "每张卡片仅承载一类基金，主动型与被动型分别输出。目标是让用户在同一行完成基金识别、收益比较与两个购买渠道额度的对照。")
    fixed_table(doc, [
        ("区块", "固定内容", "规则"),
        ("眉题", "ACTIVE / PASSIVE QDII · RETURN & QUOTA", "英文小字；随基金类型切换"),
        ("主标题", "QDII 主动型基金直销 / 代销额度", "被动型仅替换“主动型”；标题不得额外加粗"),
        ("日期", "YYYY — MM — DD", "使用当次生成日期"),
        ("列", "基金名称｜近一年收益率｜直销额度｜代销额度", "四列右侧数值对齐；无序号列"),
        ("数据源", "config.json + 现有抓取逻辑", "config.json 只定义基金池与显示名"),
    ], [1.05, 2.55, 2.9])
    add_heading(doc, "筛选、排序与缺失处理", level=2)
    for text in [
        "直销额度只接受基金公司官网当前可验证的数据；未实现、抓取失败或状态非 ok 的基金不进入卡片。",
        "已在抓取前排除的未实现直销额度基金代码：016702、017093、018036、019442。",
        "同一类别内按近一年收益率由高到低排序；收益率相同再按基金代码升序。",
        "基金代码跟随基金名称显示为次级文字；额度数值使用紧凑写法，如“1万”。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "2. 版式与信息层级")
    fixed_table(doc, [
        ("项目", "设计值 / 行为"),
        ("画布", "1080 × 1440 px，竖版小红书卡片"),
        ("外框", "四角 56 px 金色 L 形细线，距边约 40 px"),
        ("标题区", "眉题 y=112；标题 y=156；日期 y=228；中间短金线 y=282"),
        ("表头", "y=330 开始的开放式文字行；不填充、不圆角、不做独立卡片"),
        ("列锚点", "基金名称 x=82 左对齐；收益率 x=620、直销 x=840、代销 x=998 右对齐"),
        ("数据行", "60 px 行高；表头下仅一条金色细线；行间使用低透明度细分隔线"),
        ("连续性", "基金之间不留卡片间隔，不使用逐项方块容器"),
    ], [1.55, 4.95])

    doc.add_page_break()
    add_heading(doc, "3. 字体与颜色令牌")
    add_text(doc, "中文采用方案 C：宋体带来更克制的财经编辑感；数字采用衬线字体，强化收益与额度的可扫描性。")
    fixed_table(doc, [
        ("用途", "字体", "规格"),
        ("中文标题、基金名称、中文状态", "Songti SC（宋体）", "/System/Library/Fonts/Supplemental/Songti.ttc"),
        ("百分比、纯数字、日期", "DM Serif Display", "/Users/shareit/Library/Fonts/DMSerifDisplay-Regular.ttf"),
        ("英文眉题", "Cormorant Garamond", "/Users/shareit/Library/Fonts/CormorantGaramond[wght].ttf"),
    ], [1.65, 1.85, 3.0])
    add_text(doc, "标题权重规则：主标题使用常规宋体，不额外加粗；视觉重量由 52 px 字号与浅金色完成。", size=9.5, color=RGBColor(120, 75, 40), after=12)
    fixed_table(doc, [
        ("令牌", "色值", "用途"),
        ("主背景", "#05060B", "近黑深蓝底；只叠加低透明度金色柔光"),
        ("主金色", "#D4B26A", "眉题、表头、线条、角标"),
        ("浅金", "#F2E9D2", "标题、普通额度、日期"),
        ("红色", "#EF3B3B", "正收益与“暂停”状态"),
        ("正文", "#E7ECF3", "基金名称"),
        ("次级", "#8A93A6", "基金代码与辅助说明"),
        ("深蓝（保留）", "#141A2E", "旧卡片式表头/交替行的兼容色；当前方案不做表头填充"),
    ], [1.25, 1.15, 4.1])
    add_text(doc, "色彩上以深色、金色、红色形成主要视觉关系；不引入绿色，避免增加视觉负担。", size=9.5, color=RGBColor(120, 75, 40), after=0)

    add_heading(doc, "4. 状态表达与交付验收")
    fixed_table(doc, [
        ("字段", "显示规则"),
        ("近一年收益率", "正收益使用 #EF3B3B；负收益或缺失使用金色/破折号，保持三色体系。"),
        ("直销额度 / 代销额度", "“暂停”使用 #EF3B3B；其他额度使用 #F2E9D2。"),
        ("基金名称", "文本过长时省略，基金代码以 #8A93A6 紧随其后。"),
        ("表头", "必须是文字 + 单一金色分隔线，禁止恢复圆角深蓝底胶囊。"),
    ], [1.55, 4.95])
    add_heading(doc, "生成前检查清单", level=2)
    for text in [
        "主动型、被动型各输出一张 1080 × 1440 PNG；每张只包含当前直销额度已验证的基金。",
        "确认标题文字为“QDII 主动型基金直销 / 代销额度”或“QDII 被动型基金直销 / 代销额度”。",
        "确认四个列锚点不变，所有三列数值右对齐，名称列不与收益率列产生过大空档。",
        "确认未产生底部注释、序号列或单基金卡片间隔。",
        "将输出清单与被排除基金保留在 manifest 中，便于后续追溯。",
    ]:
        add_bullet(doc, text)

    doc.core_properties.title = "QDII 直销 / 代销额度卡片设计规范"
    doc.core_properties.subject = "小红书基金额度卡片的视觉与数据规则"
    doc.core_properties.author = "QDII Quota Radar"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
